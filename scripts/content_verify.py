#!/usr/bin/env python3
"""Block ungrounded source-to-outline, SVG, and PPTX content."""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
import tempfile
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path


HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
ADMIN_HEADING_RE = re.compile(
    r"^\s*(?:[-*+]\s+)?(?P<number>(?:[IVXLCDM]+|\d+(?:\.\d+)*|[A-Z]))[.)]?\s+(?P<title>.+?)\s*$",
    re.IGNORECASE,
)
SOURCE_COMMENT_RE = re.compile(r"<!--\s*source:\s*(.*?)\s*-->", re.IGNORECASE)
SOURCE_REF_RE = re.compile(
    r"(?P<file>raw_source[^\s,:]*\.md):L(?P<start>\d+)(?:-L(?P<end>\d+))?",
    re.IGNORECASE,
)
ATOMIC_RE = re.compile(
    r"(?<![\w])(?:"
    r"Q[1-4]|"
    r"\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?(?![.,]\d)|"
    r"\d+(?:[.,]\d+)*/\d+(?:[.,]\d+)*|"
    r"\d+(?:[.,]\d+)*(?:\s*(?:%|G|K|M|tỷ|triệu|nghìn|đồng|VNĐ|VND|USD|EUR|ha|km|MW|GW|GB|TB))?"
    r")(?![\w])",
    re.IGNORECASE,
)
WORD_RE = re.compile(r"\b[^\W\d_][\w-]*\b", re.UNICODE)
GENERIC_TITLE_WORDS = {
    "báo", "chỉ", "công", "dự", "giải", "hạ", "kết", "lũy", "mật", "ngày",
    "nguồn", "nhiệm", "nội", "phần", "slide", "số", "source", "tháng", "tình",
    "trạm", "tuần", "year",
}
QUALIFIERS = tuple(sorted({
    "không hoàn thành", "chưa hoàn thành", "không đạt", "chưa đạt", "hoàn thành",
    "không triển khai", "đang triển khai", "chờ phê duyệt", "đã phê duyệt",
    "đang xây dựng", "xóa nhầm", "đang thực hiện", "dự kiến", "lũy kế",
    "chậm", "mất", "đạt", "tăng", "giảm", "chờ", "tuần", "tháng", "quý", "năm",
    "not completed", "not achieved", "in progress", "increase", "decrease", "completed",
    "not deployed", "pending approval", "under construction", "approved", "delayed",
    "expected", "cumulative", "week", "month", "quarter", "year",
}, key=len, reverse=True))
MAX_SOURCE_RANGE_LINES = 30
IGNORE_FIELDS = {"layout", "visualization", "bố cục", "trực quan hóa"}
CONTENT_FIELDS = {"title", "subtitle", "info", "core message", "content", "tiêu đề", "nội dung"}
SLIDE_RE = re.compile(r"^#{3,6}\s+Slide\s+(\d+)\b", re.IGNORECASE)
SLIDE_DISPOSITION_RE = re.compile(
    r"^P\d{1,3}(?:\s*[–-]\s*P?\d{1,3})?(?:\s*,\s*P\d{1,3}(?:\s*[–-]\s*P?\d{1,3})?)*$",
    re.IGNORECASE,
)
EXCLUDED_RE = re.compile(r"^Excluded\s*[—–-]\s*\S", re.IGNORECASE)


def normalize(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).casefold()
    return re.sub(r"\s+", " ", value).strip()


def plain_text(value: str) -> str:
    value = SOURCE_COMMENT_RE.sub("", value)
    value = re.sub(r"!?\[([^]]*)\]\([^)]+\)", r"\1", value)
    value = re.sub(r"[*_`~]", "", value)
    return re.sub(r"\s+", " ", value).strip(" |\t")


def _phrase_tokens(value: str) -> list[str]:
    normalized = normalize(value)
    occupied: list[tuple[int, int]] = []
    found = []
    for phrase in QUALIFIERS:
        for match in re.finditer(rf"(?<!\w){re.escape(phrase)}(?!\w)", normalized):
            span = match.span()
            if not any(span[0] < end and start < span[1] for start, end in occupied):
                occupied.append(span)
                found.append(phrase)
    return found


def extract_tokens(value: str) -> list[str]:
    """Return locale-preserving factual atoms; never rewrite punctuation."""
    found = [match.group(0).strip() for match in ATOMIC_RE.finditer(value)]
    found.extend(_phrase_tokens(value))
    words = [match.group(0) for match in WORD_RE.finditer(value)]
    found.extend(
        word for word in words
        if (len(word) > 1 and word.isupper()) or any(character.isupper() for character in word[1:])
    )
    for index, word in enumerate(words):
        if index and word[0].isupper() and normalize(word) not in GENERIC_TITLE_WORDS:
            found.append(word)
        if index + 1 < len(words) and word[0].isupper() and words[index + 1][0].isupper():
            found.extend((word, words[index + 1]))
    unique: dict[str, str] = {}
    for token in found:
        key = normalize(token)
        if key and key not in unique:
            unique[key] = token
    return list(unique.values())


def token_keys(value: str) -> set[str]:
    return {normalize(token) for token in extract_tokens(value)}


def artifact_token_keys(value: str) -> set[str]:
    """Ignore prose-only labels while retaining factual sentences and identifiers."""
    keys: set[str] = set()
    for sentence in split_sentences(value):
        if ATOMIC_RE.search(sentence) or _phrase_tokens(sentence):
            keys |= token_keys(sentence)
            continue
        words = [match.group(0) for match in WORD_RE.finditer(sentence)]
        for index, word in enumerate(words):
            internal_cap = not word.isupper() and any(character.isupper() for character in word[1:])
            ascii_acronym = len(word) > 1 and word.isascii() and word.isupper()
            if internal_cap or ascii_acronym:
                keys.add(normalize(word))
            if (
                index + 1 < len(words)
                and word.istitle()
                and words[index + 1].istitle()
                and normalize(word) not in GENERIC_TITLE_WORDS
            ):
                keys.update((normalize(word), normalize(words[index + 1])))
    return keys


def split_sentences(value: str) -> list[str]:
    return [
        text for part in re.split(r"(?<=[.!?。！？])\s+|\n+", value)
        if (text := plain_text(part))
    ]


def extract_section(spec: str, roman: str) -> tuple[str, int, str]:
    lines = spec.splitlines()
    start = level = None
    for index, line in enumerate(lines):
        match = HEADING_RE.match(line.strip())
        if match and re.match(rf"^{roman}(?:\b|[.\s:—-])", plain_text(match.group(2)), re.IGNORECASE):
            start, level = index + 1, len(match.group(1))
            break
    if start is None or level is None:
        raise ValueError(f"design_spec.md is missing section {roman}")
    end = len(lines)
    for index in range(start, len(lines)):
        match = HEADING_RE.match(lines[index].strip())
        if match and len(match.group(1)) <= level:
            end = index
            break
    return "\n".join(lines[start:end]), start + 1, "\n".join(lines[:start - 1])


def _heading_title(line: str) -> str | None:
    match = HEADING_RE.match(line.strip())
    if match:
        return plain_text(match.group(2))
    match = ADMIN_HEADING_RE.match(line)
    if not match:
        return None
    raw = match.group("title").strip()
    title = plain_text(raw)
    letters = "".join(character for character in title if character.isalpha())
    bold = raw.startswith("**") and raw.endswith("**")
    if not bold and letters and letters != letters.upper():
        return None
    return f"{match.group('number')}. {title}"


def _coverage_key(title: str) -> str:
    title = re.sub(r"^(?:[IVXLCDM]+|\d+(?:\.\d+)*|[A-Z])[.)]?\s+", "", plain_text(title), flags=re.IGNORECASE)
    return normalize(title)


def source_sections(text: str) -> tuple[str, list[tuple[int, str, str]]]:
    """Return non-overlapping source units, including content before the first heading."""
    lines = text.splitlines()
    headings = [(index, title) for index, line in enumerate(lines) if (title := _heading_title(line))]
    if headings:
        sections: list[tuple[int, str, str]] = []
        first = headings[0][0]
        if plain_text("\n".join(lines[:first])):
            sections.append((1, "Document preamble", "\n".join(lines[:first])))
        for position, (start, title) in enumerate(headings):
            end = headings[position + 1][0] if position + 1 < len(headings) else len(lines)
            sections.append((start + 1, title, "\n".join(lines[start + 1:end])))
        return "headings", sections
    return "whole document", [(1, "Document preamble", text)]


def coverage_map(section_i: str, first_line: int) -> dict[str, list[tuple[str, int, str]]]:
    lines = section_i.splitlines()
    start = next(
        (index + 1 for index, line in enumerate(lines)
         if re.match(r"^#{3,6}\s+Source Coverage Map\s*$", line.strip(), re.IGNORECASE)),
        None,
    )
    if start is None:
        raise ValueError("design_spec.md section I is missing '### Source Coverage Map'")
    result: dict[str, list[tuple[str, int, str]]] = {}
    for offset, line in enumerate(lines[start:], start + 1):
        if line.strip().startswith("#"):
            break
        if not line.strip().startswith("|"):
            continue
        cells = [plain_text(cell) for cell in line.strip().strip("|").split("|")]
        if len(cells) < 2 or normalize(cells[0]) == "source section" or re.fullmatch(r"[-: ]+", cells[0]):
            continue
        key = _coverage_key(cells[0])
        result.setdefault(key, []).append((cells[1], first_line + offset - 1, cells[0]))
    if not result:
        raise ValueError("Source Coverage Map has no entries")
    return result


def parse_refs(comment: str) -> list[tuple[str, int, int]]:
    refs = []
    for match in SOURCE_REF_RE.finditer(comment):
        start = int(match.group("start"))
        refs.append((match.group("file"), start, int(match.group("end") or start)))
    return refs


def _source_units(filename: str, lines: list[str]) -> tuple[str, list[dict]]:
    mode, sections = source_sections("\n".join(lines))
    units = []
    for index, (start, title, _) in enumerate(sections):
        end = sections[index + 1][0] - 1 if index + 1 < len(sections) else len(lines)
        units.append({
            "id": f"{filename}:{index}",
            "start": start,
            "end": end,
            "title": title,
        })
    return mode, units


def _disposition_slides(disposition: str) -> set[int] | None:
    if not SLIDE_DISPOSITION_RE.fullmatch(disposition):
        return None
    slides: set[int] = set()
    for item in re.split(r"\s*,\s*", disposition):
        match = re.fullmatch(r"P(\d{1,3})(?:\s*[–-]\s*P?(\d{1,3}))?", item, re.IGNORECASE)
        if not match:
            return None
        start = int(match.group(1))
        end = int(match.group(2) or start)
        if end < start:
            return None
        slides.update(range(start, end + 1))
    return slides


def outline_claims(section: str, first_line: int) -> list[dict]:
    claims: list[dict] = []
    pending: list[int] = []
    slide = None
    for offset, raw in enumerate(section.splitlines()):
        stripped = raw.strip()
        slide_match = SLIDE_RE.match(stripped)
        if slide_match:
            slide = int(slide_match.group(1))
            pending = []
            continue
        comments = SOURCE_COMMENT_RE.findall(raw)
        line = SOURCE_COMMENT_RE.sub("", stripped)
        created: list[int] = []
        if line and not line.startswith("#") and not re.fullmatch(r"[-|: ]+", line):
            line = re.sub(r"^(?:>|[-*+]\s+|\d+[.)]\s+)+", "", line).strip()
            field = re.match(r"(?:\*\*)?([^:*]+)(?:\*\*)?\s*:\s*(.*)$", line)
            if field:
                label = normalize(plain_text(field.group(1)))
                if label in IGNORE_FIELDS:
                    line = ""
                elif label in CONTENT_FIELDS:
                    line = field.group(2)
            if line.endswith(":"):
                line = ""
            for sentence in split_sentences(line):
                if extract_tokens(sentence):
                    claims.append({
                        "line": first_line + offset,
                        "slide": slide,
                        "text": sentence,
                        "refs": [],
                    })
                    created.append(len(claims) - 1)
        if created:
            pending = created
        if comments and pending:
            refs = [ref for comment in comments for ref in parse_refs(comment)]
            for index in pending:
                claims[index]["refs"].extend(refs)
            pending = []
    return claims


def _evidence(refs: list[tuple[str, int, int]], sources: dict[str, list[str]]) -> tuple[str, list[str]]:
    chunks = []
    for filename, start, end in refs:
        if filename not in sources:
            raise ValueError(f"source reference uses unknown file: {filename}")
        lines = sources[filename]
        if start < 1 or end < start or end > len(lines):
            raise ValueError(f"source reference is outside {filename}:L1-L{len(lines)}: {filename}:L{start}-L{end}")
        chunks.append("\n".join(lines[start - 1:end]))
    text = "\n".join(chunks)
    windows = split_sentences(text) + [plain_text(line) for line in text.splitlines() if plain_text(line)]
    return text, windows


def _add(report: dict, severity: str, code: str, location: str, message: str) -> None:
    report[severity].append({"code": code, "location": location, "message": message})


def _verify_claims(
    claims: list[dict],
    sources: dict[str, list[str]],
    units_by_file: dict[str, list[dict]],
    report: dict,
) -> None:
    reported_invalid_refs: set[tuple[str, int, int, str]] = set()
    for claim in claims:
        location = f"design_spec.md:{claim['line']}"
        tokens = extract_tokens(claim["text"])
        claim["units"] = []
        if not claim["refs"]:
            _add(report, "errors", "MISSING_SOURCE_REF", location, claim["text"])
            continue
        invalid = False
        for filename, start, end in claim["refs"]:
            code = message = None
            if filename not in sources or start < 1 or end < start or end > len(sources.get(filename, [])):
                code = "INVALID_SOURCE_REF"
                message = f"invalid source range: {filename}:L{start}-L{end}"
            elif end - start + 1 > MAX_SOURCE_RANGE_LINES:
                code = "OVERBROAD_SOURCE_REF"
                message = f"{filename}:L{start}-L{end} exceeds {MAX_SOURCE_RANGE_LINES} lines"
            if code:
                key = (filename, start, end, code)
                if key not in reported_invalid_refs:
                    _add(report, "errors", code, location, message)
                    reported_invalid_refs.add(key)
                invalid = True
                continue
            containing = [
                unit for unit in units_by_file.get(filename, [])
                if unit["start"] <= start and end <= unit["end"]
            ]
            if len(containing) != 1:
                key = (filename, start, end, "CROSS_SECTION_SOURCE_REF")
                if key not in reported_invalid_refs:
                    _add(
                        report,
                        "errors",
                        "CROSS_SECTION_SOURCE_REF",
                        location,
                        f"{filename}:L{start}-L{end} must stay inside one Source Coverage Map unit",
                    )
                    reported_invalid_refs.add(key)
                invalid = True
                continue
            claim["units"].append(containing[0]["id"])
        if invalid:
            continue
        try:
            evidence, windows = _evidence(claim["refs"], sources)
        except ValueError as error:
            _add(report, "errors", "INVALID_SOURCE_REF", location, str(error))
            continue
        evidence_keys = token_keys(evidence)
        absent = [token for token in tokens if normalize(token) not in evidence_keys]
        if absent:
            _add(report, "errors", "UNSUPPORTED_TOKEN", location, f"{absent} | {claim['text']}")
            continue
        keys = {normalize(token) for token in tokens}
        if len(keys) > 1 and not any(keys <= token_keys(window) for window in windows):
            _add(report, "errors", "UNSUPPORTED_PAIRING", location, claim["text"])


def _verify_coverage(units: list[dict], claims: list[dict], report: dict) -> None:
    used: dict[str, set[int]] = {unit["id"]: set() for unit in units}
    indexed = {unit["id"]: unit for unit in units}
    for claim in claims:
        slide = claim.get("slide")
        if slide is None:
            continue
        for unit_id in set(claim.get("units", [])):
            unit = indexed[unit_id]
            disposition = unit.get("disposition", "")
            location = f"design_spec.md:{claim['line']}"
            if EXCLUDED_RE.match(disposition):
                _add(
                    report,
                    "errors",
                    "EXCLUDED_SECTION_CITED",
                    location,
                    f"Slide {slide:02d} cites excluded section: {unit['title']}",
                )
                continue
            slides = _disposition_slides(disposition)
            if slides is not None and slide not in slides:
                _add(
                    report,
                    "errors",
                    "COVERAGE_SLIDE_MISMATCH",
                    location,
                    f"{unit['title']} maps to {disposition}, not P{slide:02d}",
                )
                continue
            used[unit_id].add(slide)
    for unit in units:
        slides = _disposition_slides(unit.get("disposition", ""))
        if slides is None:
            continue
        missing = slides - used[unit["id"]]
        if missing:
            _add(
                report,
                "errors",
                "COVERAGE_SLIDE_UNUSED",
                f"design_spec.md:{unit['map_line']}",
                f"{unit['title']} has no anchored claim for "
                + ", ".join(f"P{slide:02d}" for slide in sorted(missing)),
            )


def _parse_artifact(text: str) -> dict[int, str]:
    slides: dict[int, list[str]] = {}
    current = None
    for line in text.splitlines():
        match = re.match(r"^## Slide (\d+)\s*$", line.strip(), re.IGNORECASE)
        if match:
            current = int(match.group(1))
            slides[current] = []
        elif current is not None and not line.lstrip().startswith("!["):
            slides[current].append(line)
    return {number: "\n".join(lines) for number, lines in slides.items()}


def _artifact_text(artifact: Path) -> dict[int, str]:
    if artifact.is_dir():
        svg_files = sorted(artifact.glob("*.svg"))
        if not svg_files:
            raise ValueError(f"artifact directory has no SVG files: {artifact}")
        slides = {}
        for number, svg_file in enumerate(svg_files, 1):
            try:
                root = ET.parse(svg_file).getroot()
            except ET.ParseError as error:
                raise ValueError(f"invalid SVG {svg_file}: {error}") from error
            blocks = []
            for element in root.iter():
                if element.tag.rsplit("}", 1)[-1] == "text":
                    if text := plain_text(" ".join(element.itertext())):
                        blocks.append(text)
            slides[number] = "\n".join(blocks)
        return slides
    if not artifact.is_file() or artifact.suffix.lower() != ".pptx":
        raise ValueError(f"artifact must be an SVG directory or existing .pptx: {artifact}")
    try:
        from source_to_md.ppt_to_md import convert_presentation_to_markdown
    except ImportError as error:
        raise ValueError(f"cannot load ppt_to_md.py: {error}") from error
    with tempfile.TemporaryDirectory(prefix="content_verify_pptx_") as directory:
        output = Path(directory) / "artifact.md"
        markdown = convert_presentation_to_markdown(str(artifact), str(output))
        if not markdown:
            raise ValueError(f"could not extract text from {artifact}")
        return _parse_artifact(markdown)


def _verify_artifact(
    artifact: Path,
    claims: list[dict],
    report: dict,
) -> None:
    slides = _artifact_text(artifact)
    by_slide: dict[int, list[dict]] = {}
    for claim in claims:
        if claim["slide"] is not None:
            by_slide.setdefault(claim["slide"], []).append(claim)
    for number, slide_text in slides.items():
        if number not in by_slide:
            unsupported = sorted(
                artifact_token_keys(slide_text) - {str(number), f"{number:02d}", f"p{number:02d}"}
            )
            if unsupported:
                _add(
                    report,
                    "errors",
                    "ARTIFACT_SLIDE_WITHOUT_SOURCES",
                    f"artifact:Slide {number}",
                    str(unsupported),
                )
    for number, slide_claims in by_slide.items():
        location = f"artifact:Slide {number}"
        if number not in slides:
            _add(report, "errors", "MISSING_SLIDE", location, "slide exists in section IX but not in PPTX")
            continue
        slide_text = slides[number]
        artifact_keys = artifact_token_keys(slide_text)
        all_artifact_keys = token_keys(slide_text)
        spec_keys = token_keys("\n".join(claim["text"] for claim in slide_claims))
        valid_windows: list[set[str]] = [token_keys(claim["text"]) for claim in slide_claims]
        allowed = spec_keys | {str(number), f"{number:02d}", f"p{number:02d}"}
        unsupported = sorted(artifact_keys - allowed)
        if unsupported:
            _add(report, "errors", "ARTIFACT_UNSUPPORTED_TOKEN", location, str(unsupported))
        for claim in slide_claims:
            missing = [token for token in extract_tokens(claim["text"]) if normalize(token) not in all_artifact_keys]
            if missing:
                _add(report, "warnings", "ARTIFACT_MISSING_CLAIM", location, f"{missing} | {claim['text']}")
        for sentence in split_sentences(slide_text):
            keys = token_keys(sentence) - {str(number)}
            qualifier_keys = {normalize(token) for token in _phrase_tokens(sentence)}
            if len(keys) > 1 and (keys & qualifier_keys or ATOMIC_RE.search(sentence)):
                if keys <= allowed and not any(keys <= window for window in valid_windows):
                    _add(report, "errors", "ARTIFACT_UNSUPPORTED_PAIRING", location, sentence)


def verify(project: Path, artifact: Path | None = None) -> dict:
    spec_path = project / "design_spec.md"
    source_paths = sorted((project / "sources").glob("raw_source*.md"))
    if not spec_path.is_file():
        raise ValueError(f"missing {spec_path}")
    if not source_paths:
        raise ValueError(f"no raw_source*.md found in {project / 'sources'}")

    spec = spec_path.read_text(encoding="utf-8")
    section_i, section_i_line, _ = extract_section(spec, "I")
    section_ix, first_line, _ = extract_section(spec, "IX")
    mappings = coverage_map(section_i, section_i_line)
    sources = {path.name: path.read_text(encoding="utf-8").splitlines() for path in source_paths}
    report = {"errors": [], "warnings": [], "modes": {}}

    used: dict[str, int] = {}
    units: list[dict] = []
    units_by_file: dict[str, list[dict]] = {}
    for path in source_paths:
        mode, source_units = _source_units(path.name, sources[path.name])
        report["modes"][path.name] = mode
        units_by_file[path.name] = source_units
        units.extend(source_units)
        for unit in source_units:
            key = _coverage_key(unit["title"])
            location = f"{path.name}:{unit['start']}"
            position = used.get(key, 0)
            if key not in mappings or position >= len(mappings[key]):
                _add(report, "errors", "UNMAPPED_SOURCE_SECTION", location, unit["title"])
                continue
            disposition, map_line, map_title = mappings[key][position]
            used[key] = position + 1
            unit["disposition"] = disposition
            unit["map_line"] = map_line
            if not (SLIDE_DISPOSITION_RE.fullmatch(disposition) or EXCLUDED_RE.match(disposition)):
                _add(report, "errors", "INVALID_DISPOSITION", f"design_spec.md:{map_line}", f"{map_title} | {disposition}")
    for key, entries in mappings.items():
        for _, line, title in entries[used.get(key, 0):]:
            _add(report, "errors", "UNKNOWN_COVERAGE_SECTION", f"design_spec.md:{line}", title)

    claims = outline_claims(section_ix, first_line)
    _verify_claims(claims, sources, units_by_file, report)
    range_errors = {"INVALID_SOURCE_REF", "OVERBROAD_SOURCE_REF", "CROSS_SECTION_SOURCE_REF"}
    if not any(item["code"] in range_errors for item in report["errors"]):
        _verify_coverage(units, claims, report)
    if artifact is not None:
        _verify_artifact(artifact, claims, report)
    return report


def print_report(report: dict) -> None:
    for source, mode in report["modes"].items():
        print(f"SOURCE MODE | {source} | {mode}")
    for severity in ("errors", "warnings"):
        for finding in report[severity]:
            print(f"{severity[:-1].upper()} | {finding['code']} | {finding['location']} | {finding['message']}")
    print("SUMMARY | " + " | ".join(f"{key}={len(report[key])}" for key in ("errors", "warnings")))


def self_test() -> None:
    with tempfile.TemporaryDirectory() as directory:
        project = Path(directory)
        (project / "sources").mkdir()
        source_path = project / "sources" / "raw_source.md"
        source_text = (
            "Mật độ báo cáo\n\n1. **TÌNH HÌNH TRONG TUẦN**\n"
            "MyViettel tăng 1.100 khách hàng.\n"
            "KPI đạt 127/361 trong Q3.\n"
            "Sự cố khác ghi nhận 33.999 thuê bao.\n"
        )
        source_path.write_text(source_text, encoding="utf-8")
        clean = (
            "## I. Project Information\n\n### Source Coverage Map\n\n"
            "| Source section | Disposition |\n| --- | --- |\n"
            "| Document preamble | Excluded — administrative |\n"
            "| TÌNH HÌNH TRONG TUẦN | P02 |\n\n"
            "## IX. Content Outline\n\n#### Slide 02\n"
            "- **Content**:\n  - MyViettel tăng 1.100 khách hàng.\n"
            "    <!-- source: raw_source.md:L4-L4 -->\n"
            "  - KPI đạt 127/361 trong Q3.\n"
            "    <!-- source: raw_source.md:L5-L5 -->\n\n"
            "## X. Technical Constraints\n"
        )
        spec = project / "design_spec.md"
        spec.write_text(clean, encoding="utf-8")
        report = verify(project)
        assert not report["errors"], report

        source_path.write_text(source_text + "\n".join(f"Dòng {number}" for number in range(25)), encoding="utf-8")
        spec.write_text(clean.replace("raw_source.md:L4-L4", "raw_source.md:L1-L31"), encoding="utf-8")
        report = verify(project)
        assert any(item["code"] == "OVERBROAD_SOURCE_REF" for item in report["errors"]), report

        source_path.write_text(source_text, encoding="utf-8")
        spec.write_text(clean.replace("raw_source.md:L4-L4", "raw_source.md:L2-L3"), encoding="utf-8")
        report = verify(project)
        assert any(item["code"] == "CROSS_SECTION_SOURCE_REF" for item in report["errors"]), report

        spec.write_text(clean.replace("| TÌNH HÌNH TRONG TUẦN | P02 |", "| TÌNH HÌNH TRONG TUẦN | P02–P03 |"), encoding="utf-8")
        report = verify(project)
        assert any(item["code"] == "COVERAGE_SLIDE_UNUSED" for item in report["errors"]), report

        source_path.write_text(source_text.replace("Mật độ báo cáo", "Báo cáo Q3"), encoding="utf-8")
        excluded_claim = clean.replace(
            "#### Slide 02\n",
            "#### Slide 02\n- **Content**:\n  - Q3.\n    <!-- source: raw_source.md:L1-L1 -->\n",
        )
        spec.write_text(excluded_claim, encoding="utf-8")
        report = verify(project)
        assert any(item["code"] == "EXCLUDED_SECTION_CITED" for item in report["errors"]), report

        source_path.write_text(source_text, encoding="utf-8")
        spec.write_text(clean, encoding="utf-8")

        svg_dir = project / "svg_output"
        svg_dir.mkdir()
        (svg_dir / "01.svg").write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" viewBox="0 0 1920 1080"/>',
            encoding="utf-8",
        )
        svg_slide = svg_dir / "02.svg"
        svg_slide.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" viewBox="0 0 1920 1080">'
            '<text x="100" y="50">HẠNG MỤC</text>'
            '<text x="100" y="100">MyViettel tăng 1.100 khách hàng.</text>'
            '<text x="100" y="200">KPI đạt 127/361 trong Q3.</text><text x="1800" y="1000">02</text></svg>',
            encoding="utf-8",
        )
        report = verify(project, svg_dir)
        assert not report["errors"] and not report["warnings"], report

        exporter = Path(__file__).with_name("svg_to_pptx.py")
        final_output = project / "final.pptx"
        result = subprocess.run(
            [sys.executable, str(exporter), str(project), "-o", str(final_output), "-q", "-t", "none", "-a", "none"],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 0 and final_output.is_file(), result.stdout + result.stderr

        svg_slide.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg"><text>Tăng khách hàng MyViettel: 1.100.</text>'
            '<text>KPI đạt 127/361 trong Q3.</text></svg>',
            encoding="utf-8",
        )
        report = verify(project, svg_dir)
        assert not report["errors"] and not report["warnings"], report

        svg_slide.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg"><text>MyViettel tăng 1.100 khách hàng.</text>'
            '<text>KPI đạt 127/361 trong Q3.</text><text>Sự cố khác: 33.999 thuê bao.</text>'
            '<text>NetBuy</text></svg>',
            encoding="utf-8",
        )
        report = verify(project, svg_dir)
        unsupported = [item for item in report["errors"] if item["code"] == "ARTIFACT_UNSUPPORTED_TOKEN"]
        assert unsupported and "netbuy" in unsupported[0]["message"], report
        failed_output = project / "failed.pptx"
        result = subprocess.run(
            [sys.executable, str(exporter), str(project), "-o", str(failed_output), "-q", "-t", "none", "-a", "none"],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 1 and not failed_output.exists(), result.stdout + result.stderr
        assert not list(project.glob("pptx_content_gate_*"))

        from pptx import Presentation
        from pptx.util import Inches

        def write_pptx(path: Path, text: str, extra_slide: str | None = None) -> None:
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
            box.text = text
            if extra_slide:
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
                box.text = extra_slide
            presentation.save(path)

        artifact = project / "clean.pptx"
        write_pptx(artifact, "MyViettel tăng 1.100 khách hàng.\nKPI đạt 127/361 trong Q3.")
        report = verify(project, artifact)
        assert not report["errors"] and not report["warnings"], report

        write_pptx(artifact, "MyViettel tăng 9.999 khách hàng.\nKPI đạt 127/361 trong Q3.")
        report = verify(project, artifact)
        assert any(item["code"] == "ARTIFACT_UNSUPPORTED_TOKEN" for item in report["errors"]), report

        write_pptx(artifact, "Công ty Sao Mai tăng 1.100 khách hàng.\nKPI đạt 127/361 trong Q3.")
        report = verify(project, artifact)
        assert any(item["code"] == "ARTIFACT_UNSUPPORTED_TOKEN" for item in report["errors"]), report

        write_pptx(artifact, "MyViettel tăng 1.100 khách hàng.\nKPI không đạt 127/361 trong Q3.")
        report = verify(project, artifact)
        assert any(item["code"] == "ARTIFACT_UNSUPPORTED_TOKEN" for item in report["errors"]), report

        write_pptx(artifact, "MyViettel tăng 127/361 trong Q3.\nKPI đạt 1.100 khách hàng.")
        report = verify(project, artifact)
        assert any(item["code"] == "ARTIFACT_UNSUPPORTED_PAIRING" for item in report["errors"]), report

        write_pptx(artifact, "MyViettel tăng 1.100 khách hàng.\nKPI đạt 127/361 trong Q3.", "Số mới 8.888")
        report = verify(project, artifact)
        assert any(item["code"] == "ARTIFACT_SLIDE_WITHOUT_SOURCES" for item in report["errors"]), report

        write_pptx(artifact, "MyViettel tăng 1.100 khách hàng.")
        report = verify(project, artifact)
        assert any(item["code"] == "ARTIFACT_MISSING_CLAIM" for item in report["warnings"]), report

        spec.write_text(clean.replace("127/361", "999/361"), encoding="utf-8")
        assert any(item["code"] == "UNSUPPORTED_TOKEN" for item in verify(project)["errors"])

        bad_pairing = clean.replace("MyViettel tăng 1.100 khách hàng.", "MyViettel tăng 33.999 khách hàng.")
        bad_pairing = bad_pairing.replace("raw_source.md:L4-L4", "raw_source.md:L4-L6")
        spec.write_text(bad_pairing, encoding="utf-8")
        assert any(item["code"] == "UNSUPPORTED_PAIRING" for item in verify(project)["errors"])

        spec.write_text(clean.replace("KPI đạt 127/361", "KPI không đạt 127/361"), encoding="utf-8")
        assert any(item["code"] == "UNSUPPORTED_TOKEN" for item in verify(project)["errors"])

        spec.write_text(clean.replace("    <!-- source: raw_source.md:L5-L5 -->\n", ""), encoding="utf-8")
        assert any(item["code"] == "MISSING_SOURCE_REF" for item in verify(project)["errors"])
    print("content_verify self-test: OK")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project", nargs="?", type=Path, help="project containing design_spec.md and sources/")
    parser.add_argument("--artifact", type=Path, help="SVG directory or final PPTX to verify")
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        self_test()
        return 0
    if args.project is None:
        parser.error("project is required unless --self-test is used")
    try:
        report = verify(args.project.resolve(), args.artifact.resolve() if args.artifact else None)
        print_report(report)
    except (OSError, UnicodeError, ValueError) as error:
        print(f"ERROR | {error}")
        return 2
    return 1 if report["errors"] or report["warnings"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
