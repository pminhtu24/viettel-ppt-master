#!/usr/bin/env python3
"""Run with python3 scripts/test_content_grounding.py; synthetic sources only."""
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from content_check import (chart_check, calculate, parse_source_number, validate,
                           export_gate, snapshot, review_entries, svg_check, check_page)
from source_index import build_index, read_blocks, sha256

ROOT = Path(__file__).resolve().parents[1]
CATALOG = json.loads((ROOT / 'templates/charts/charts_index.json').read_text())['charts']


def fact(key, value=10, **changes):
    return dict(id=key, value=value, raw=str(value), locale='en', text=f'{value} items',
                entity='orders', unit='items', period='2026-01-01', scope='all', status='observed',
                category=key, series='A', **changes)


def visualization(key, values):
    roles = CATALOG[key]['data_requirements']['roles']
    return dict(id='v1', template=key, observations=[dict(zip(roles, values))],
                evidence={c:dict(refs=values, reason='Synthetic source establishes this relation')
                          for c in CATALOG[key]['data_requirements']['semantic_checks']})


class GroundingTests(unittest.TestCase):
    def test_numeric_locale_and_safe_calculation(self):
        self.assertEqual(parse_source_number('1.234,5', 'vi'), 1234.5)
        self.assertEqual(parse_source_number('1,234.5', 'en'), 1234.5)
        for raw in ('', '-', 'NaN', '1.23.4'):
            with self.assertRaises(ValueError):
                parse_source_number(raw, 'vi')
        values = {'a': fact('a', 132), 'b': fact('b', 101)}
        item = dict(id='r', inputs=['a','b'], operation='remainder', value=31)
        self.assertEqual(calculate(item, values), 31)
        item['value'] = 28
        with self.assertRaises(ValueError): calculate(item, values)
        item['operation'] = '__import__("os")'
        with self.assertRaises(ValueError): calculate(item, values)
        item.update(operation='percent', value=130.7)
        self.assertGreater(calculate(item, values), 100)

    def test_every_catalog_entry_requires_real_inputs_and_evidence(self):
        # Every one of the 71 templates rejects an empty observation set.
        for key in CATALOG:
            with self.subTest(key=key):
                viz = visualization(key, ['a'])
                viz['observations'] = []
                with self.assertRaises(ValueError): chart_check(viz, {'a': fact('a')}, CATALOG)
        for key in ('kpi_cards','bar_chart','horizontal_bar_chart','heatmap_chart','radar_chart',
                    'process_flow','matrix_2x2','timeline','roadmap_vertical'):
            viz = visualization(key, ['a'])
            chart_check(viz, {'a':fact('a')}, CATALOG)
            viz['observations'][0]['value'] = 'fabricated'
            with self.assertRaises(ValueError): chart_check(viz, {'a':fact('a')}, CATALOG)

    def test_time_no_synthetic_history(self):
        for key in ('line_chart','area_chart','stacked_area_chart','dual_axis_line_chart'):
            vals = {'a':fact('a', 79), 'b':fact('b', 81)}
            vals['b']['period'] = '2026-01-08'
            viz = visualization(key, ['a']); viz['observations'].append({'value':'b'})
            if key == 'stacked_area_chart':
                vals['total_a'] = dict(vals['a'], id='total_a')
                vals['total_b'] = dict(vals['b'], id='total_b')
                viz['totals'] = [dict(whole='total_a', parts=['a']), dict(whole='total_b', parts=['b'])]
            chart_check(viz, vals, CATALOG)
            viz['observations'].pop()
            with self.assertRaises(ValueError): chart_check(viz, vals, CATALOG)

    def test_compositions_and_pareto(self):
        for key in ('pie_chart','donut_chart','stacked_bar_chart','treemap_chart','pareto_chart'):
            vals = {'a':fact('a',82),'b':fact('b',17),'total':fact('total',99)}
            viz = visualization(key,['a']); viz['observations'].append({'value':'b'})
            viz['totals']=[dict(whole='total',parts=['a','b'])]
            chart_check(viz, vals, CATALOG)
            vals['b']['value']=18
            with self.assertRaises(ValueError): chart_check(viz, vals, CATALOG)

    def test_progress_above_target_and_mixed_units(self):
        for key in ('progress_bar_chart','gauge_chart','bullet_chart'):
            vals={'a':fact('a',120),'b':fact('b',100)}
            viz=visualization(key,['a']); viz['observations'][0]['target']='b'; chart_check(viz,vals,CATALOG)
            vals['b']['unit']='square meters'
            with self.assertRaises(ValueError): chart_check(viz,vals,CATALOG)

    def test_funnel_waterfall_and_flow(self):
        vals={'a':fact('a',100),'b':fact('b',60),'c':fact('c',40)}
        viz=visualization('funnel_chart',['a']); viz['observations'] += [{'value':'b'},{'value':'c'}]
        chart_check(viz,vals,CATALOG)
        del viz['evidence']['same-cohort']
        with self.assertRaises(ValueError): chart_check(viz,vals,CATALOG)
        vals['a'].update(value=100,status='start'); vals['b'].update(value=-20,status='change'); vals['c'].update(value=80,status='end')
        viz=visualization('waterfall_chart',['a']); viz['observations'] += [{'value':'b'},{'value':'c'}]
        chart_check(viz,vals,CATALOG)
        vals['c']['value']=85
        with self.assertRaises(ValueError): chart_check(viz,vals,CATALOG)
        vals={'a':fact('a',10),'b':fact('b',10)}
        vals['a'].update(entity='A',category='B'); vals['b'].update(entity='B',category='C')
        viz=visualization('sankey_chart',['a']); viz['observations'].append({'value':'b'})
        chart_check(viz,vals,CATALOG)
        vals['b']['value']=8
        with self.assertRaises(ValueError): chart_check(viz,vals,CATALOG)

    def test_distribution_xy_schedule_pairs(self):
        for key in ('scatter_chart','bubble_chart','box_plot_chart','gantt_chart'):
            roles=CATALOG[key]['data_requirements']['roles']
            vals={r:fact(r,i+1) for i,r in enumerate(roles)}
            if key=='gantt_chart':
                vals['start']['value']='2026-01-01'; vals['end']['value']='2026-02-01'
            viz=visualization(key,list(vals)); chart_check(viz,vals,CATALOG)
            del viz['observations'][0][roles[-1]]
            with self.assertRaises(ValueError): chart_check(viz,vals,CATALOG)
        for key in ('grouped_bar_chart','dumbbell_chart','butterfly_chart'):
            vals={'a':fact('a'), 'b':fact('b')}
            vals['b'].update(category='a',series='B')
            viz=visualization(key,['a']); viz['observations'].append({'value':'b'})
            chart_check(viz,vals,CATALOG)
            vals['b']['series']=''
            with self.assertRaises(ValueError): chart_check(viz,vals,CATALOG)

    def test_end_to_end_gate_and_stale_review(self):
        with tempfile.TemporaryDirectory() as temp:
            project=Path(temp); (project/'sources').mkdir(); (project/'svg_output').mkdir()
            source=project/'sources/report.md'; source.write_text('Store: 82 orders fulfilled.\n\n| Item | Count |\n| -- | -- |\n| A | 82 |\n')
            first=build_index(project); self.assertEqual(first,build_index(project))
            blocks=read_blocks(project); ref=next(iter(blocks))
            f=fact('f1',82); f['sources']=[dict(ref=ref,quote=blocks[ref])]
            page=dict(page='P01',svg='01.svg',source_index_sha256=sha256(project/'sources/source_index.json'),
                      facts=[f],derived=[],conclusions=[],unresolved=[],visualizations=[visualization('kpi_cards',['f1'])],required=['f1'])
            def save():
                (project/'design_spec.md').write_text('content_contract: 1\n```content-lock\n'+json.dumps(page)+'\n```\n')
            save(); validate(project,'plan')
            svg=project/'svg_output/01.svg'
            valid='<svg xmlns="http://www.w3.org/2000/svg" data-content-contract="1"><g data-chart="v1"><text data-claims="f1">82 items</text></g></svg>'
            svg.write_text(valid); validate(project,'svg')
            self.assertFalse(export_gate([svg])) # review absent
            entries={'P01':dict(status='pass',reviewer='test',notes='Synthetic fixture; no business conclusion',
                **{key:True for key in ['source_meaning','numbers_units','scope_period_status','titles_conclusions','required_content','visual_encoding']},charts={'v1':'82 matches source'})}
            review_entries(entries,[page])
            (project/'content_review.json').write_text(json.dumps(dict(version=1,hashes=snapshot(project,[page]),pages=entries)))
            self.assertTrue(export_gate([svg]))
            svg.write_text(valid.replace('82 items','100 items'))
            with self.assertRaises(ValueError): validate(project,'svg')
            self.assertFalse(export_gate([svg]))
            svg.write_text(valid+'\n'); self.assertFalse(export_gate([svg]))
            svg.write_text(valid); source.write_text(source.read_text()+'New source line\n')
            with self.assertRaises(ValueError): validate(project,'plan')
            build_index(project)
            with self.assertRaises(ValueError): validate(project,'plan')
            page['source_index_sha256']=sha256(project/'sources/source_index.json'); save()
            page['facts'][0]['sources'][0]['quote']='invented completion'; save()
            with self.assertRaises(ValueError): validate(project,'plan')

    def test_new_project_cannot_fall_back_to_legacy(self):
        from project_manager import ProjectManager
        with tempfile.TemporaryDirectory() as temp:
            project = Path(ProjectManager(temp).init_project('grounding', brand_profile='custom_override'))
            with self.assertRaisesRegex(ValueError, 'cannot export as legacy'):
                validate(project, 'plan')
            (project/'design_spec.md').write_text('# Outline without a contract')
            with self.assertRaises(ValueError): validate(project, 'plan')

    def test_hidden_required_values_and_geometry(self):
        values = {'a':fact('a',5), 'b':fact('b',15)}
        viz = visualization('progress_bar_chart',['a'])
        viz['observations'][0]['target'] = 'b'
        viz['geometry'] = [dict(element='bar', attribute='width', fact='a', domain=[0,15], range=[0,600])]
        page = dict(required=['a'], visualizations=[viz])
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'slide.svg'
            valid='<svg data-content-contract="1"><g data-chart="v1"><rect id="bar" width="200"/><text data-claims="a">5 items</text></g></svg>'
            path.write_text(valid); svg_check(path,page,values)
            for bad in [valid.replace('width="200"','width="500"'),
                        valid.replace('>5 items<','>items<'),
                        valid.replace('<text ', '<text display="none" ')]:
                path.write_text(bad)
                with self.assertRaises(ValueError): svg_check(path,page,values)

    def test_source_substrings_conflicts_and_review_completeness(self):
        f=fact('a',1); f['sources']=[dict(ref='s',quote='There are 101 orders')]
        page=dict(facts=[f],derived=[],conclusions=[],unresolved=[],visualizations=[],required=['a'])
        with self.assertRaises(ValueError): check_page(page,{'s':'There are 101 orders'},CATALOG)
        f.update(value=101,raw='101',text='101 orders')
        check_page(page,{'s':'There are 101 orders'},CATALOG)
        page['unresolved']=[dict(description='Conflicting main denominator',material=True,resolution='flagged')]
        with self.assertRaises(ValueError): check_page(page,{'s':'There are 101 orders'},CATALOG)
        f['status']='unverified'
        with self.assertRaises(ValueError): chart_check(visualization('kpi_cards',['a']),{'a':f},CATALOG)
        with self.assertRaises(ValueError): review_entries({'P01':dict(status='pass')},[dict(page='P01',visualizations=[])])

    def test_percent_without_target_and_duplicate_parts(self):
        f=fact('a',125); f['unit']='%'
        chart_check(visualization('progress_bar_chart',['a']),{'a':f},CATALOG)
        values={'a':fact('a',82),'b':fact('b',17),'total':fact('total',99)}
        viz=visualization('pie_chart',['a']); viz['observations'] += [dict(value='b'),dict(value='total')]
        viz['totals']=[dict(whole='total',parts=['a','b'])]
        with self.assertRaises(ValueError): chart_check(viz,values,CATALOG)
        viz['observations']=[dict(value='a'),dict(value='a'),dict(value='b')]
        with self.assertRaises(ValueError): chart_check(viz,values,CATALOG)

    def test_word_tables_images_and_legacy_failure(self):
        from docx import Document
        from source_to_md.doc_to_md import convert_to_markdown, _convert_legacy_doc
        import base64
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp); doc=Document(); doc.add_heading('Nguồn thử nghiệm',0)
            doc.add_paragraph('Giữ phạm vi ngoài dự án; dự kiến hoàn thành.')
            table=doc.add_table(rows=3,cols=2)
            for i,row in enumerate(table.rows):
                row.cells[0].text=f'Hạng mục {i}'; row.cells[1].text='Chưa có số liệu' if i==2 else str(i)
            merged=doc.add_table(rows=2,cols=2); merged.cell(0,0).merge(merged.cell(0,1)).text='Ô gộp'
            image=folder/'pixel.png'; image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScLbtAAAAABJRU5ErkJggg=='))
            doc.add_picture(str(image)); doc.save(folder/'source.docx')
            text=convert_to_markdown(str(folder/'source.docx'),str(folder/'source.md'))
            self.assertIn('| Hạng mục 2 | Chưa có số liệu |',text)
            self.assertIn('colspan="2"',text); self.assertIn('ngoài dự án',text)
            self.assertTrue((folder/'source_files/image_manifest.json').exists())
            with patch.dict(os.environ,{'LIBREOFFICE_BIN':''}), patch('shutil.which',return_value=None):
                with self.assertRaisesRegex(ValueError,'requires LibreOffice'):
                    _convert_legacy_doc(folder/'missing.doc',folder/'fail.md')
            with patch.dict(os.environ,{'LIBREOFFICE_BIN':'false'}):
                with self.assertRaises(ValueError): _convert_legacy_doc(folder/'missing.doc',folder/'fail.md')


if __name__=='__main__':
    unittest.main()
